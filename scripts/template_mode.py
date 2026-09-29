# -*- coding: utf-8 -*-
"""템플릿 모드 — 사용자가 가진 매뉴얼 pptx 템플릿의 형태 그대로 매뉴얼을 짓는다.

템플릿의 **크롬**(표지·목차·본문 머리/꼬리·글꼴·색)은 템플릿에서 가져오고, **본문**(캡처
검은 테두리·균일 폭·긴 캡처 분할·표 자동 분할·배지 배분·검증 게이트)은 번들 빌더가
그대로 그린다. 서식을 통째로 복제하던 방식에서 디자인이 깨지던 문제를 이 분리로 막는다.

구성
  analyze(path)          템플릿을 읽어 매니페스트(레이아웃 역할·채울 자리·서식·본문 영역·
                         잔존 감시 문자열)를 만든다. 휴리스틱이므로 처음 한 번은 검토한다.
  load_manifest(path)    사용자 로컬(~/.claude/manual-gen/templates/)에서 파일 해시로 찾고,
                         없으면 분석해 저장한다. 템플릿을 고치면 해시가 바뀌어 다시 분석한다.
                         (개인 경로·표본 문구가 담기므로 skill 저장소 밖에 둔다)
  TemplateDeck           템플릿 사본을 열어 표본 슬라이드를 지우고, 레이아웃 사본의 고정
                         텍스트를 이번 매뉴얼 문구로 바꾼 뒤, 레이아웃으로 새 슬라이드를 만든다.
  leftover_check(pptx)   산출물 전체(슬라이드·레이아웃·마스터)에 템플릿 표본 문구가 남았는지.
  등록부                 register·resolve_template — 한 번 쓴 템플릿을 이름으로 기억하고 기본
                         템플릿을 정한다(~/.claude/manual-gen/templates.json, 관리: template_registry.py).

템플릿 원본 파일은 절대 수정하지 않는다 — 읽기만 하고 결과는 다른 경로에 저장한다.
"""

import copy
import hashlib
import json
import math
import os
import re
import zipfile

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.oxml.ns import qn

E = 914400
STORE_DIR = os.path.join(os.path.expanduser("~"), ".claude", "manual-gen", "templates")
REGISTRY = os.path.join(os.path.expanduser("~"), ".claude", "manual-gen", "templates.json")
MANIFEST_VERSION = 2          # 형식이 바뀌면 올린다 — 저장본이 옛 형식이면 자동으로 다시 분석
NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
_GROUP, _LINE = 6, 9
_TITLE_TYPES = (PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.TITLE)
_TOC_WORD = re.compile(r"^(목차|차례|CONTENTS|Contents|contents)$")
_CHAPTER_TITLE = re.compile(r"^(\d{1,2})\.\s+\S")            # "01. 시스템 개요"
_SECTION_TITLE = re.compile(r"^(\d+(?:\.\d+)+)(\s+)(\S.*)$")   # "1.1 사이트 이용 흐름도"
_CHAPTER_LABEL = re.compile(r"^CHAPTER\s*\d+", re.I)


# ---------- 공통 도우미 ------------------------------------------------------

def file_sha1(path):
    h = hashlib.sha1()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def _in(v):
    return round(v / E, 3) if v is not None else None


def iter_text_shapes(shapes):
    """글자를 담을 수 있는 도형을 그룹 안까지 훑는다."""
    for sh in shapes:
        if sh.shape_type == _GROUP:
            yield from iter_text_shapes(sh.shapes)
        elif getattr(sh, "has_text_frame", False):
            yield sh


def _text(sh):
    return sh.text_frame.text.replace("\x0b", "\n").strip()


def _line1(t):
    return (t or "").split("\n")[0].strip()


def _ph(sh):
    return sh.placeholder_format.idx if sh.is_placeholder else None


def style_of(el):
    """도형 안 글자 서식(글꼴·크기·색·굵기). 우선순위는 렌더와 같게 — 런 > 단락 끝 >
    목록 기본값. (문서 순서로 훑으면 lstStyle 이 런보다 앞에 있어 순서가 뒤집힌다)"""
    runs = [r for r in el.iter(qn("a:rPr"))]
    ends = [r for r in el.iter(qn("a:endParaRPr"))]
    defs = [r for r in el.iter(qn("a:defRPr"))]
    font = size = color = bold = None
    for rpr in runs + ends + defs:
        ea, la = rpr.find(qn("a:ea")), rpr.find(qn("a:latin"))
        f = (ea.get("typeface") if ea is not None else None) or (la.get("typeface") if la is not None else None)
        if font is None and f and not f.startswith("+"):
            font = f
        if size is None and rpr.get("sz"):
            size = int(rpr.get("sz")) / 100
        if color is None:
            c = rpr.find("a:solidFill/a:srgbClr", NS)
            if c is not None:
                color = c.get("val").upper()
        if bold is None and rpr.get("b") is not None:
            bold = rpr.get("b") == "1"
    return {"font": font, "size": size, "color": color, "bold": bool(bold)}


def _text_left(sh):
    """글자가 실제로 시작하는 x(in) — 상자 왼쪽 + 왼쪽 여백(lIns, 기본 0.1in) + 첫 단락 들여쓰기(marL)."""
    body = sh.text_frame._txBody
    bp = body.find(qn("a:bodyPr"))
    lins = bp.get("lIns") if bp is not None else None
    p = body.find(qn("a:p"))
    ppr = p.find(qn("a:pPr")) if p is not None else None
    marl = ppr.get("marL") if ppr is not None else None
    if marl is None:
        lvl = body.find("a:lstStyle/a:lvl1pPr", NS)
        marl = lvl.get("marL") if lvl is not None else None
    return (sh.left + int(lins if lins is not None else 91440) + int(marl or 0)) / E


def _align(sh):
    """첫 단락 정렬(l·ctr·r) — 단락 > 목록 기본값 순."""
    body = sh.text_frame._txBody
    p = body.find(qn("a:p"))
    ppr = p.find(qn("a:pPr")) if p is not None else None
    a = ppr.get("algn") if ppr is not None else None
    if a is None:
        lvl = body.find("a:lstStyle/a:lvl1pPr", NS)
        a = lvl.get("algn") if lvl is not None else None
    return a or "l"


def _nlines(sh):
    return max(1, len(_text(sh).split("\n")))


def _pitch(boxes, size):
    """목차 한 줄 간격(in) — 표본 상자 높이 ÷ 줄 수 중 가장 촘촘한 값. 글자에 맞추지 않고
    길게 늘여 둔 자리(레이아웃의 마지막 자리 등)는 글자 크기 대비 범위를 벗어나 버린다."""
    lo, hi = size / 72 * 1.6, size / 72 * 3.2
    vals = [_in(b.height) / _nlines(b) for b in boxes if b.height]
    vals = [v for v in vals if lo <= v <= hi]
    return round(min(vals), 3) if vals else None


def _depth(t):
    """절 번호 깊이 — '1.1 …' 1, '3.1.1 …' 2, 절이 아니면 0."""
    m = _SECTION_TITLE.match(t or "")
    return m.group(1).count(".") if m else 0


_R_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
_RT_OK = ("/image", "/hyperlink")     # 다른 쪽으로 옮겨도 관계를 다시 걸 수 있는 종류


def _has_text(el):
    return any((t.text or "").strip() for t in el.iter(qn("a:t")))


def _decor(slide, avoid=None):
    """표본 슬라이드에만 그려진 장식(글자 없는 선·도형·그림) — 템플릿을 연 사람이 보는 모습은
    이것까지 포함이므로 새 쪽마다 같이 얹는다. 본문·목차 영역(avoid: x0,y0,x1,y1 in)과 겹치는
    것(본문을 가림)과 그림·링크 외의 관계(차트·SmartArt 등)를 가진 것은 뺀다.
    반환: (shape_id 목록, 뺀 개수)"""
    ids, skipped = [], 0
    for sh in slide.shapes:
        if sh.is_placeholder or _has_text(sh._element):
            continue
        x0, y0 = _in(sh.left or 0), _in(sh.top or 0)
        x1, y1 = x0 + _in(sh.width or 0), y0 + _in(sh.height or 0)
        overlaps = avoid is not None and x0 < avoid[2] and x1 > avoid[0] and y0 < avoid[3] and y1 > avoid[1]
        rids = [v for x in sh._element.iter() for k, v in x.attrib.items() if k.startswith(_R_NS)]
        movable = all(r in slide.part.rels and slide.part.rels[r].reltype.endswith(_RT_OK) for r in rids)
        if overlaps or not movable:
            skipped += 1
            continue
        ids.append(sh.shape_id)
    return ids, skipped


def _hlines(layout):
    """레이아웃의 가로 선 — [(y, x, w, 색)], 위에서 아래 순."""
    out = []
    for sh in layout.shapes:
        if sh.shape_type == _LINE or (sh.height == 0 and sh.width and sh.width > E):
            c = sh._element.find(".//a:ln/a:solidFill/a:srgbClr", NS)
            out.append((_in(sh.top), _in(sh.left), _in(sh.width), c.get("val").upper() if c is not None else None))
    return sorted(out)


def _same(a, b):
    return a is not None and b is not None and a.part is b.part


def _layout_ref(prs, layout):
    for mi, m in enumerate(prs.slide_masters):
        for li, lay in enumerate(m.slide_layouts):
            if lay.part is layout.part:
                return {"master": mi, "index": li, "name": lay.name}
    return None


def set_text_keep_style(txBody, text):
    """글상자 서식은 그대로 두고 글자만 바꾼다 — 첫 단락·첫 런만 남기고 그 런의 글자를
    교체한다. 런이 없으면(빈 안내문) 단락 끝 서식을 새 런에 옮겨 쓴다."""
    ps = txBody.findall(qn("a:p"))
    if not ps:
        p0 = txBody.makeelement(qn("a:p"), {})
        txBody.append(p0)
        ps = [p0]
    p0 = ps[0]
    for p in ps[1:]:
        txBody.remove(p)
    for child in list(p0):
        if child.tag in (qn("a:br"), qn("a:fld")):
            p0.remove(child)
    runs = p0.findall(qn("a:r"))
    if runs:
        for r in runs[1:]:
            p0.remove(r)
        t = runs[0].find(qn("a:t"))
        t.text = text
        return
    r = p0.makeelement(qn("a:r"), {})
    end = p0.find(qn("a:endParaRPr"))
    rpr = copy.deepcopy(end) if end is not None else p0.makeelement(qn("a:rPr"), {})
    rpr.tag = qn("a:rPr")
    r.append(rpr)
    t = r.makeelement(qn("a:t"), {})
    t.text = text
    r.append(t)
    if end is not None:
        end.addprevious(r)
    else:
        p0.append(r)


# ---------- 한 줄 맞춤 --------------------------------------------------------
# 템플릿의 글자 자리(표지 제목·매뉴얼 구분·대상 알약·러닝헤더 등)는 표본 글자 길이에 맞춰
# 크기가 정해져 있다. 이번 문구가 더 길면 줄이 꺾여 옆 요소와 겹치므로, 실제 글꼴로 폭을
# 재어 한 줄에 들게 한다 — 배경 도형이 있는 알약은 넓히고, 나머지는 글자를 줄인다.

_FONT_DIRS = [os.path.join(os.environ.get("LOCALAPPDATA", ""), "Microsoft", "Windows", "Fonts"),
              os.path.join(os.environ.get("WINDIR", r"C:\Windows"), "Fonts")]
_font_index = {}              # 글꼴 이름(모든 언어) → 파일
_font_regular = {}            # 그 파일이 보통 굵기인지
_scanned = []                 # 훑은 폴더
MIN_FIT = 0.6                 # 글자는 원래 크기의 60% 까지만 줄인다


def _sfnt_names(path):
    """글꼴 파일(TTF·OTF·TTC 첫 글꼴) 이름표의 family·full name — 모든 언어(한글 이름 포함).
    반환: (이름 집합, 보통 굵기 여부) — 같은 family 를 굵기 파일 여럿이 나눠 쓰면 보통 굵기를 고른다."""
    import struct
    names, subs = set(), set()
    try:
        with open(path, "rb") as f:
            head = f.read(12)
            if head[:4] == b"ttcf":
                f.seek(12)
                f.seek(struct.unpack(">I", f.read(4))[0])
                head = f.read(12)
            num = struct.unpack(">H", head[4:6])[0]
            recs = f.read(16 * num)
            data = None
            for i in range(num):
                tag, _c, off, ln = struct.unpack(">4sIII", recs[16 * i:16 * i + 16])
                if tag == b"name":
                    f.seek(off)
                    data = f.read(ln)
                    break
        if data is None:
            return names, False
        count, soff = struct.unpack(">HH", data[2:6])
        for i in range(count):
            pid, _e, _l, nid, ln, off = struct.unpack(">HHHHHH", data[6 + 12 * i:18 + 12 * i])
            if pid == 3 and nid in (1, 2, 4):  # 16(굵기 공통 이름)은 넣지 않는다 — 굵기 파일이 섞인다
                try:
                    v = data[soff + off:soff + off + ln].decode("utf-16-be").strip()
                except UnicodeDecodeError:
                    continue
                (subs if nid == 2 else names).add(v)
    except (OSError, struct.error, IndexError):
        pass
    return names, bool(subs & {"Regular", "Normal", "Book", "Roman"})


def font_file(name):
    """글꼴 이름 → 파일 경로(없으면 None). 사용자 글꼴 폴더부터 훑고, 못 찾으면 시스템 폴더."""
    if not name:
        return None
    for d in _FONT_DIRS:
        if name in _font_index:
            break
        if d in _scanned or not os.path.isdir(d):
            continue
        _scanned.append(d)
        for fn in os.listdir(d):
            if fn.lower().endswith((".ttf", ".otf", ".ttc")):
                p = os.path.join(d, fn)
                names, regular = _sfnt_names(p)
                for n in names:
                    if n not in _font_index or (regular and not _font_regular.get(n)):
                        _font_index[n], _font_regular[n] = p, regular
    return _font_index.get(name)


def _is_wide(c):
    o = ord(c)
    return 0x1100 <= o <= 0x11FF or 0x3130 <= o <= 0x318F or 0xAC00 <= o <= 0xD7A3 or 0x4E00 <= o <= 0x9FFF


def text_width(text, size_pt, font=None, spc=0):
    """한 줄 글자 폭(in) — 글꼴 파일이 있으면 실측(Pillow), 없으면 글자 종류별 추정(넉넉히)."""
    w = None
    path = font_file(font)
    if path:
        try:
            from PIL import ImageFont
            w = ImageFont.truetype(path, 200).getlength(text) / 200 * size_pt / 72
        except Exception:
            w = None
    if w is None:
        em = sum(0.95 if _is_wide(c) else 0.3 if c == " " else 0.64 if c.isupper() else 0.56 for c in text)
        w = em * size_pt / 72
    return w + len(text) * (spc or 0) / 100 / 72


def _run_props(*sources):
    """첫 런의 크기(pt)·글꼴·자간 — 런 > 목록 기본값(lvl1) 순, 앞의 출처가 우선
    (슬라이드 자리표시자 → 레이아웃 자리표시자 → 마스터 자리표시자·스타일의 a:defRPr)."""
    size = font = spc = None
    for src in sources:
        if src is None:
            continue
        cands = (src,) if src.tag == qn("a:defRPr") else (
            src.find(".//a:r/a:rPr", NS), src.find("a:lstStyle/a:lvl1pPr/a:defRPr", NS))
        for rpr in cands:
            if rpr is None:
                continue
            if size is None and rpr.get("sz"):
                size = int(rpr.get("sz")) / 100
            if spc is None and rpr.get("spc"):
                spc = int(rpr.get("spc"))
            if font is None:
                for tag in ("a:ea", "a:latin"):
                    x = rpr.find(tag, NS)
                    if x is not None and x.get("typeface") and not x.get("typeface").startswith("+"):
                        font = x.get("typeface")
                        break
    return size, font, spc or 0


def _xfrm(el):
    """도형의 a:xfrm (그룹이면 grpSpPr, 아니면 spPr)."""
    return el.find("p:grpSpPr/a:xfrm", NS) if el.tag == qn("p:grpSp") else el.find("p:spPr/a:xfrm", NS)


def _box(el):
    x = _xfrm(el)
    if x is None or x.find("a:off", NS) is None:
        return None
    o, e = x.find("a:off", NS), x.find("a:ext", NS)
    return int(o.get("x")), int(o.get("y")), int(e.get("cx")), int(e.get("cy"))


def _widen(el, delta, delta_ch=None):
    """오른쪽으로 넓힌다 — 그룹은 쪽 좌표 크기(ext)와 자식 좌표 크기(chExt)를 함께."""
    e = _xfrm(el).find("a:ext", NS)
    e.set("cx", str(int(e.get("cx")) + delta))
    ch = _xfrm(el).find("a:chExt", NS)
    if ch is not None:
        ch.set("cx", str(int(ch.get("cx")) + (delta if delta_ch is None else delta_ch)))


def _insets(*txbodies):
    left = right = None
    for tb in txbodies:
        bp = tb.find("a:bodyPr", NS) if tb is not None else None
        if bp is None:
            continue
        left = bp.get("lIns") if left is None else left
        right = bp.get("rIns") if right is None else right
    return int(left if left is not None else 91440) + int(right if right is not None else 91440)


def _pill_bg(el, box, tree):
    """글상자를 품은 배경(채움 있고 글자 없는 도형) — 같은 트리(그룹·레이아웃)에서 찾는다."""
    tol = int(E * 0.06)
    for sib in tree.iterchildren(qn("p:sp")):
        if sib is el or _has_text(sib) or sib.find("p:spPr/a:solidFill", NS) is None:
            continue
        b = _box(sib)
        if b and b[0] - tol <= box[0] and b[1] - tol <= box[1] and box[0] + box[2] <= b[0] + b[2] + tol \
                and box[1] + box[3] <= b[1] + b[3] + tol:
            return sib
    return None


def master_defaults(lay_el, master_el):
    """레이아웃 자리표시자가 물려받는 마스터 서식 — 같은 종류의 마스터 자리표시자 글상자와
    제목·본문 스타일(txStyles)의 1단계 기본값. 레이아웃에 글자 크기가 없는 자리(제목 등)용."""
    if lay_el is None or master_el is None:
        return ()
    ph = lay_el.find(".//p:nvPr/p:ph", NS)
    typ = ph.get("type", "body") if ph is not None else "body"
    want = "title" if typ in ("title", "ctrTitle") else typ
    out = []
    for sp in master_el.iter(qn("p:sp")):
        mph = sp.find(".//p:nvPr/p:ph", NS)
        if mph is not None and mph.get("type", "body") == want:
            out.append(sp.find(qn("p:txBody")))
            break
    style = "p:titleStyle" if want == "title" else "p:bodyStyle"
    d = master_el.find(f"p:txStyles/{style}/a:lvl1pPr/a:defRPr", NS)
    out.append(d)
    return tuple(x for x in out if x is not None)


def fit_line(el, slide_w, lay_el=None, lay_tree=None, inherit=()):
    """글상자 el 의 글자가 한 줄에 들도록 — 알약(배경 도형)이면 넓히고, 아니면 글자를 줄인다.
    el 이 슬라이드 자리표시자면 lay_el(레이아웃 자리표시자)에서 크기·여백·위치를 물려받고
    (그래도 없는 크기는 inherit — 마스터 서식), 배경은 lay_tree(레이아웃 도형 트리)에서 찾는다.
    반환: 경고 문구 또는 None."""
    tb = el.find(qn("p:txBody"))
    ltb = lay_el.find(qn("p:txBody")) if lay_el is not None else None
    text = _all_text(el)
    size, font, spc = _run_props(tb, ltb, *inherit)
    own = _box(el)
    box = own or (_box(lay_el) if lay_el is not None else None)
    if not text or size is None or box is None:
        return None
    parent = el.getparent()
    scale = 1.0
    if parent.tag == qn("p:grpSp"):
        gx = _xfrm(parent)
        ext, chext = gx.find("a:ext", NS), gx.find("a:chExt", NS)
        if ext is not None and chext is not None and int(chext.get("cx")):
            scale = int(ext.get("cx")) / int(chext.get("cx"))
    avail = (box[2] * scale - _insets(tb, ltb)) / E
    need = text_width(text, size, font, spc)
    if need <= avail * 0.985:
        return None
    # 알약 — 배경과 함께 오른쪽으로 넓힌다(쪽 오른쪽 여백 0.3in 안에서)
    tree = parent if parent.tag == qn("p:grpSp") else (lay_tree if own is None and lay_tree is not None else parent)
    src = el if own is not None else lay_el
    bg = _pill_bg(src, box, tree)
    geom = src.find("p:spPr/a:prstGeom", NS)
    self_pill = (src.find("p:spPr/a:solidFill", NS) is not None and geom is not None
                 and geom.get("prst") == "roundRect")          # 글상자 자신이 알약인 경우
    if bg is not None or self_pill:
        delta = int((need - avail + 0.08) * E / scale)          # 자식 좌표계 기준
        grp = parent if parent.tag == qn("p:grpSp") else None
        right = (_box(grp)[0] + _box(grp)[2] + delta * scale) if grp is not None else box[0] + box[2] + delta
        if right <= slide_w - E * 0.3:
            if own is None:                      # 슬라이드 자리표시자 — 위치를 물려받아 자기 xfrm 로
                sp_pr = el.find(qn("p:spPr"))
                x = sp_pr.makeelement(qn("a:xfrm"), {})
                o = x.makeelement(qn("a:off"), {"x": str(box[0]), "y": str(box[1])})
                e = x.makeelement(qn("a:ext"), {"cx": str(box[2] + delta), "cy": str(box[3])})
                x.append(o)
                x.append(e)
                sp_pr.insert(0, x)
            else:
                _widen(el, delta)
            if bg is not None:
                _widen(bg, delta)
            if grp is not None:
                _widen(grp, int(delta * scale), delta)          # 그룹 크기는 쪽 좌표, chExt 는 자식 좌표
            return None
    # 글자를 줄인다 — 반 포인트 단위, 원래 크기의 MIN_FIT 까지
    new = max(size * MIN_FIT, math.floor(size * avail * 0.97 / need * 2) / 2)
    for rpr in list(tb.iter(qn("a:rPr"))) + list(tb.iter(qn("a:endParaRPr"))):
        rpr.set("sz", str(int(round(new * 100))))
    if text_width(text, new, font, spc) > avail * 1.01:
        return f"'{text[:24]}' — 자리({avail:.2f}in)보다 길어 {new:g}pt 로 줄여도 줄이 꺾일 수 있음"
    return None


# ---------- 분석 -------------------------------------------------------------

def _roles(prs):
    slides = list(prs.slides)
    layouts = [lay for m in prs.slide_masters for lay in m.slide_layouts]

    def texts_of(obj):
        return [_line1(_text(sh)) for sh in iter_text_shapes(obj.shapes)]

    cover = slides[0].slide_layout if slides else next(
        (lay for lay in layouts if any(p.placeholder_format.type == PP_PLACEHOLDER.CENTER_TITLE
                                       for p in lay.placeholders)), None)

    def is_toc(lay):
        ts = texts_of(lay) + [t for s in slides if _same(s.slide_layout, lay) for t in texts_of(s)]
        return any(_TOC_WORD.match(t) for t in ts)

    toc = next((s.slide_layout for s in slides if not _same(s.slide_layout, cover) and is_toc(s.slide_layout)),
               None) or next((lay for lay in layouts if not _same(lay, cover) and is_toc(lay)), None)

    def is_content_slide(s):
        if _same(s.slide_layout, cover) or _same(s.slide_layout, toc):
            return False
        return any(_CHAPTER_TITLE.match(t) for t in texts_of(s))

    content_slide = next((s for s in slides if is_content_slide(s)), None)
    content = content_slide.slide_layout if content_slide else next(
        (s.slide_layout for s in reversed(slides)
         if not _same(s.slide_layout, cover) and not _same(s.slide_layout, toc)), None)
    return cover, toc, content, content_slide


def _cover_spec(cover, cover_slide):
    """표지의 채울 자리 — 자리표시자는 표본 슬라이드에 채워진 값, 없으면 안내문으로 성격을 본다."""
    filled = {}
    if cover_slide is not None:
        for ph in cover_slide.placeholders:
            filled[ph.placeholder_format.idx] = _text(ph)
    slots, guard, keep = [], [], []
    for sh in iter_text_shapes(cover.shapes):
        idx = _ph(sh)
        prompt = _text(sh)
        sample = filled.get(idx, "") if idx is not None else ""
        value = _line1(sample or prompt)
        if idx is not None and sh.placeholder_format.type in _TITLE_TYPES:
            field = "system"
            src = sample or ("" if "넣어" in prompt else prompt)
            if _line1(src):
                guard.append(_line1(src))
        elif re.fullmatch(r"(19|20)\d{2}", value):
            field = "year"
        elif re.fullmatch(r"\S{1,8}용", value):
            field = "audience"
        elif re.search(r"(매뉴얼|메뉴얼)", value) and not re.search(r"[A-Za-z]", value):
            field = "kind"
        elif idx is not None and (re.search(r"[A-Za-z]{3,}", value) or "영문" in prompt):
            field = "meta"                   # 영문 부제 자리 — 우리에겐 영문명이 없으므로 버전·날짜
            if re.search(r"[A-Za-z]{3,}", _line1(sample)):
                guard.append(_line1(sample))
        elif idx is not None:
            field = "clear"
        else:
            field = "keep"
            if prompt:
                keep.append(prompt)
            continue
        slot = {"field": field, "sample": value[:60]}
        if idx is not None:
            slot["ph"] = idx
        else:
            slot["static"] = prompt
        slots.append(slot)
    style = next((style_of(sh._element) for sh in iter_text_shapes(cover.shapes)
                  if sh.is_placeholder and sh.placeholder_format.type in _TITLE_TYPES), {})
    return {"slots": slots, "keep_static": keep}, guard, style.get("color")


def _chrome_spec(layout):
    """목차·본문 레이아웃의 머리·꼬리 자리 — 위 가는 선보다 위는 러닝헤더, 아래 선보다
    아래는 쪽 번호·장 표기·라벨. 선이 없으면 쪽 가장자리 0.6in 를 기준으로 삼는다."""
    lines = _hlines(layout)
    h = _in(layout.part.package.presentation_part.presentation.slide_height)
    top = min((y for y, *_ in lines if y < 1.6), default=0.6)
    bot = max((y for y, *_ in lines if y > h - 1.6), default=h - 0.6)
    slots = []
    for sh in iter_text_shapes(layout.shapes):
        idx, t = _ph(sh), _line1(_text(sh))
        y, hh = _in(sh.top), _in(sh.height)
        s = {"ph": idx} if idx is not None else {"static": _text(sh)}
        if y + hh <= top + 0.05 and t:
            s.update(role="run_head", sample=t)
        elif y >= bot - 0.02:
            if re.fullmatch(r"\d{1,3}", t):
                s.update(role="page", pad=len(t))
            elif _CHAPTER_LABEL.match(t):
                m = re.match(r"^(CHAPTER\s*)(\d+)(\.?\s*)(.*)$", t, re.I)
                num = "{n:02d}" if m and m.group(2).startswith("0") else "{n}"
                s.update(role="footer_chapter",
                         fmt=(m.group(1) + num + m.group(3) + "{title}") if m else "CHAPTER {n}. {title}")
            elif t:
                s.update(role="footer_label", text=t)
            else:
                continue
        elif idx is not None and sh.placeholder_format.type in _TITLE_TYPES and _TOC_WORD.match(t):
            s.update(role="toc_title", text=t)
        else:
            continue
        slots.append(s)
    color = next((c for *_, c in lines if c), None)
    return slots, top, bot, color


def _all_text(el):
    return " ".join((t.text or "").strip() for t in el.iter(qn("a:t"))).strip()


def _toc_gap(labels, rows, pitch_of):
    """장 블록 사이 간격(in) — 첫 장 라벨과 둘째 장 라벨 사이 마지막 줄의 아래에서 둘째 라벨까지."""
    labels = sorted(labels, key=lambda s: s.top)
    if len(labels) < 2:
        return None
    a, b = labels[0], labels[1]
    between = [r for r in rows if a.top < r.top < b.top]
    if not between:
        return None
    last = max(between, key=lambda r: r.top)
    g = round(_in(b.top) - (_in(last.top) + _nlines(last) * pitch_of(last)), 3)
    return g if 0.05 <= g <= 0.6 else None


def _toc_spec(toc, toc_slide):
    """목차 항목 서식·위치·간격 — 자리표시자(안내문)에서 장 라벨·장 제목·절·하위 절·쪽 번호의
    글꼴·크기·색을 읽고, 항목 영역(좌·폭·시작 높이)과 줄 간격을 잡는다. 레이아웃에 없는 자리
    (하위 절 등)는 표본 목차 슬라이드의 글상자에서 보충한다."""
    phs = [sh for sh in toc.placeholders]
    tx = {id(sh): _line1(_text(sh)) for sh in phs}
    label = min((sh for sh in phs if _CHAPTER_LABEL.match(tx[id(sh)])), key=lambda s: s.top, default=None)
    if label is None:
        return None
    top_of = lambda s: s.top
    foot = toc.part.package.presentation_part.presentation.slide_height - E * 1.2   # 꼬리 영역 경계

    def near(sh, other, dy=0.08):
        return abs(_in(sh.top) - _in(other.top)) <= dy

    # 높이 0 에 가까운 자리표시자(선 모양 SmartArt 자리 등)는 글자 자리가 아니다
    title = min((sh for sh in phs if sh.height > E * 0.1 and label.left - E * 0.1 <= sh.left <= label.left + E * 0.1
                 and label.top < sh.top <= label.top + E * 0.45 and tx[id(sh)]
                 and not _SECTION_TITLE.match(tx[id(sh)]) and not re.fullmatch(r"[\d\n]+", tx[id(sh)])),
                key=top_of, default=None)
    # 표본 슬라이드의 맨 위 글상자(그룹 안은 좌표계가 달라 제외) — 레이아웃에 없는 자리 보충용
    sample = [sh for sh in (toc_slide.shapes if toc_slide is not None else [])
              if getattr(sh, "has_text_frame", False) and not sh.is_placeholder and _text(sh)]
    stx = {id(sh): _line1(_text(sh)) for sh in sample}
    secs = [sh for sh in phs if _depth(tx[id(sh)]) == 1]
    subs = [sh for sh in phs if _depth(tx[id(sh)]) >= 2]
    s_secs = [sh for sh in sample if _depth(stx[id(sh)]) == 1]
    s_subs = [sh for sh in sample if _depth(stx[id(sh)]) >= 2]
    sec = min(secs, key=top_of, default=None) or min(s_secs, key=top_of, default=None)
    sub = min(subs, key=top_of, default=None) or min(s_subs, key=top_of, default=None)
    digits = lambda sh: re.fullmatch(r"\d{1,3}(\n\d{1,3})*", _text(sh) or "")
    pages_body = [sh for sh in phs if digits(sh) and sh.left > label.left + E * 2 and sh.top < foot]
    s_pages = [sh for sh in sample if digits(sh) and sh.left > label.left + E * 2 and sh.top < foot]

    def page_for(row):
        if row is None:
            return None
        pool = pages_body if any(row is p for p in phs) else s_pages
        return next((p for p in pool if near(p, row)), None)

    ch_page, sec_page, sub_page = page_for(title), page_for(sec), page_for(sub)
    if ch_page is None and title is not None:
        # 장 쪽 번호 자리가 없는 템플릿 — 표본 목차 슬라이드에서 장 제목과 같은 줄의 숫자를 찾는다
        ch_row = next((r for r in sample if stx[id(r)] == tx[id(title)]), None) or next(
            (r for r in sample if r.left <= label.left + E * 0.1 and not _depth(stx[id(r)]) and not digits(r)
             and not _CHAPTER_LABEL.match(stx[id(r)]) and r.top > label.top), None)
        if ch_row is not None:
            ch_page = next((p for p in s_pages if near(p, ch_row)), None)
    rule = next((sh for sh in phs if sh.height < E * 0.02 and sh.width > E and sh.top >= label.top - E * 0.05),
                None)
    right = max((p.left + p.width for p in pages_body), default=label.left + E * 4.7)
    num_w = max((p.width for p in pages_body), default=E * 0.45)
    st = lambda sh, fb=None: style_of(sh._element) if sh is not None else (fb or {})
    sec_st = st(sec, {"size": 9})
    sub_st = st(sub, sec_st)

    # 간격 — 레이아웃 자리의 실제 간격(라벨→장 제목→첫 절)과 표본 상자의 줄 간격
    spacing = {}
    if title is not None:
        d = _in(title.top - label.top)
        if 0.08 <= d <= 0.5:
            spacing["label_h"] = round(d, 3)
        nxt = min((s for s in secs if s.top > title.top), key=top_of, default=None)
        if nxt is not None and 0.15 <= _in(nxt.top - title.top) <= 0.8:
            spacing["title_h"] = round(_in(nxt.top - title.top), 3)
    sec_size = sec_st.get("size") or 9
    sub_size = sub_st.get("size") or sec_size
    spacing["sec_h"] = _pitch(secs + s_secs, sec_size)
    spacing["sub_h"] = _pitch(subs + s_subs, sub_size)

    def pitch_of(r):
        deep = _depth(_line1(_text(r))) >= 2
        return (spacing["sub_h"] if deep else spacing["sec_h"]) or (sub_size if deep else sec_size) / 72 * 1.85

    s_labels = [sh for sh in (toc_slide.shapes if toc_slide is not None else [])
                if _CHAPTER_LABEL.match(_all_text(sh._element))]
    spacing["gap"] = (_toc_gap([sh for sh in phs if _CHAPTER_LABEL.match(tx[id(sh)])], secs + subs, pitch_of)
                      or _toc_gap(s_labels, s_secs + s_subs, pitch_of))
    indent = round(_text_left(sub) - _text_left(sec), 3) if (sub is not None and sec is not None) else 0
    # 글자 시작 위치를 템플릿과 맞춘다 — 빌더 글상자의 왼쪽 여백(0.02in)만큼 당겨 둔다
    x = round(_text_left(label) - 0.02, 3)
    return {
        "x": x, "w": round(_in(right) - x, 3), "top": _in(label.top), "num_w": _in(num_w),
        "rule_dy": _in(rule.top - label.top) if rule is not None else 0.08,
        "page_align": _align(sec_page or ch_page) if (sec_page or ch_page) is not None else "r",
        "spacing": {k: v for k, v in spacing.items() if v is not None},
        "styles": {
            "label": st(label), "title": st(title, {"size": 14}),
            "title_page": st(ch_page) if ch_page is not None else dict(sec_st),
            "sec": sec_st, "sec_page": st(sec_page, sec_st),
            "sub": sub_st, "sub_page": st(sub_page, sub_st),
        },
        "indent": indent if indent > 0.02 else 0.2,
    }


def _content_spec(content, content_slide, top_rule, bot_rule):
    """본문 쪽 — 표본 슬라이드의 장·절 제목 상자를 복제 원형으로 삼고, 그 아래를 본문 영역으로."""
    chap = sec = None
    if content_slide is not None:
        for sh in iter_text_shapes(content_slide.shapes):
            t = _line1(_text(sh))
            if chap is None and _CHAPTER_TITLE.match(t):
                chap = sh
            elif sec is None and _SECTION_TITLE.match(t):
                sec = sh
    spec = {"body_bottom": round(bot_rule - 0.10, 3)}
    boxes = [b for b in (chap, sec) if b is not None]
    spec["body_top"] = round(max(_in(b.top + b.height) for b in boxes) + 0.14, 3) if boxes \
        else round(top_rule + 0.95, 3)
    ref = sec or chap
    if ref is not None:
        lins = ref.text_frame._txBody.find("a:bodyPr", NS).get("lIns")
        spec["body_x"] = round(_in(ref.left) + (int(lins) / E if lins else 0.1) - 0.02, 3)
    proto = {}
    for key, sh in (("chapter", chap), ("section", sec)):
        if sh is None:
            continue
        xml = sh._element
        # 관계(r:id)를 가진 도형은 다른 슬라이드로 옮기면 깨진다 — 복제하지 않는다
        if any(k.endswith("}id") or k.endswith("}embed") for el in xml.iter() for k in el.attrib):
            continue
        t = _line1(_text(sh))
        if key == "chapter":
            fmt = "{num}. {title}"
        else:
            m = _SECTION_TITLE.match(t)
            fmt = "{num}" + (m.group(2) if m else " ") + "{title}"
        proto[key] = {"shape_id": sh.shape_id, "fmt": fmt, "style": style_of(xml),
                      "box": [_in(sh.left), _in(sh.top), _in(sh.width), _in(sh.height)]}
    spec["proto"] = proto
    return spec


def analyze(path):
    prs = Presentation(path)
    w, h = _in(prs.slide_width), _in(prs.slide_height)
    if abs(w - 7.5) < 0.06 and abs(h - 10.833) < 0.06:
        orient = "portrait"
    elif abs(w - 13.333) < 0.06 and abs(h - 7.5) < 0.06:
        orient = "landscape"
    else:
        raise ValueError(f"지원하지 않는 슬라이드 크기 {w}x{h}in — 세로 A4(7.5x10.83) 또는 가로 16:9"
                         "(13.33x7.5) 템플릿만 쓸 수 있습니다")
    cover, toc, content, content_slide = _roles(prs)
    if cover is None or content is None:
        raise ValueError("표지·본문 레이아웃을 찾지 못했습니다 — 표지(첫 슬라이드)와 본문 표본 슬라이드"
                         "('01. 장 제목' 글자가 있는 쪽)가 들어 있는 템플릿이어야 합니다")
    slides = list(prs.slides)
    cover_slide = slides[0] if slides and _same(slides[0].slide_layout, cover) else None
    toc_slide = next((s for s in slides if toc is not None and _same(s.slide_layout, toc)), None)

    cover_spec, guard, cover_color = _cover_spec(cover, cover_slide)
    content_chrome, top_rule, bot_rule, rule_color = _chrome_spec(content)
    content_spec = _content_spec(content, content_slide, top_rule, bot_rule)
    content_spec["chrome"] = content_chrome
    if content_slide is not None:
        content_spec["proto_slide"] = slides.index(content_slide)
    toc_spec = None
    if toc is not None:
        toc_chrome, _t, _b, _c = _chrome_spec(toc)
        toc_spec = _toc_spec(toc, toc_slide) or {}
        toc_spec["chrome"] = toc_chrome

    # 표본 슬라이드에만 있는 장식 — 표지는 전부, 목차·본문은 항목·본문 영역과 겹치지 않는 것만
    decor = {}
    bx = content_spec.get("body_x", 0.5)
    avoid = {"cover": None,
             "content": (bx, content_spec["body_top"], w - bx, content_spec["body_bottom"])}
    if toc_spec and toc_spec.get("x") is not None:
        avoid["toc"] = (toc_spec["x"] - 0.1, toc_spec["top"] - 0.1, toc_spec["x"] + toc_spec["w"] + 0.1,
                        content_spec["body_bottom"])
    for role, sl in (("cover", cover_slide), ("toc", toc_slide), ("content", content_slide)):
        if sl is not None and role in avoid:
            ids, skipped = _decor(sl, avoid[role])
            if ids or skipped:
                decor[role] = {"slide": slides.index(sl), "ids": ids, "skipped": skipped}

    chap_style = (content_spec.get("proto", {}).get("chapter") or {}).get("style", {})
    fonts = set()
    for sh in list(iter_text_shapes(cover.shapes)) + list(iter_text_shapes(content.shapes)) + \
            (list(iter_text_shapes(toc.shapes)) if toc is not None else []):
        f = style_of(sh._element).get("font")
        if f:
            fonts.add(f)
    for spec in (toc_spec or {}).get("styles", {}).values():
        if spec.get("font"):
            fonts.add(spec["font"])
    return {
        "version": MANIFEST_VERSION,
        "template": {"path": os.path.abspath(path), "name": os.path.basename(path)},
        "size": {"w": w, "h": h, "orientation": orient},
        "layouts": {"cover": _layout_ref(prs, cover), "toc": _layout_ref(prs, toc) if toc is not None else None,
                    "content": _layout_ref(prs, content), "divider": None},
        "cover": cover_spec,
        "toc": toc_spec,
        "content": content_spec,
        "colors": {"accent": cover_color or rule_color or "2F5BD2",
                   "dark": chap_style.get("color") or "1B2A4A",
                   "rule": rule_color},
        "fonts": sorted(fonts),
        "decor": decor,
        "guard": sorted({g for g in guard if g and len(g) >= 3}),
    }


# ---------- 매니페스트 저장소 -------------------------------------------------

def manifest_path(template_path, sha):
    stem = os.path.splitext(os.path.basename(template_path))[0]
    return os.path.join(STORE_DIR, f"{stem}-{sha[:10]}.json")


def load_manifest(template_path, refresh=False):
    """(매니페스트, 새로 분석했는지, 저장 경로) — 같은 내용의 템플릿이면 저장본을 쓴다."""
    if not os.path.exists(template_path):
        raise FileNotFoundError(f"템플릿 없음: {template_path}")
    sha = file_sha1(template_path)
    mp = manifest_path(template_path, sha)
    if os.path.exists(mp) and not refresh:
        with open(mp, encoding="utf-8") as f:
            m = json.load(f)
        if m.get("version") == MANIFEST_VERSION:
            m["template"]["path"] = os.path.abspath(template_path)   # 템플릿을 옮겨도 해시로 찾는다
            return m, False, mp
    m = analyze(template_path)
    m["template"]["sha1"] = sha
    os.makedirs(STORE_DIR, exist_ok=True)
    with open(mp, "w", encoding="utf-8") as f:
        json.dump(m, f, ensure_ascii=False, indent=2)
    return m, True, mp


# ---------- 템플릿 등록부 ------------------------------------------------------
# 한 번 쓴 템플릿(회사·조직 지정 템플릿)을 이름으로 기억해, 다음 실행의 확인 표에 기본값으로
# 제시한다. 경로가 담기므로 skill 저장소 밖(사용자 폴더)에 둔다. 등록돼 있어도 적용은 매번
# 사용자 확인을 거친다 — 모르는 사이 서식이 바뀌는 것을 막기 위함(SKILL.md 1-1 의 9).

DEFAULT_REFS = ("", "default", "기본")


def load_registry():
    try:
        with open(REGISTRY, encoding="utf-8") as f:
            r = json.load(f)
    except (OSError, ValueError):
        r = {}
    r.setdefault("version", 1)
    r.setdefault("default", None)
    r.setdefault("templates", [])
    return r


def save_registry(r):
    os.makedirs(os.path.dirname(REGISTRY), exist_ok=True)
    tmp = REGISTRY + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(r, f, ensure_ascii=False, indent=2)
    os.replace(tmp, REGISTRY)


def _same_path(a, b):
    return os.path.normcase(os.path.abspath(a)) == os.path.normcase(os.path.abspath(b))


def register(path, name=None, make_default=False):
    """템플릿을 등록한다 — 같은 이름이면 경로를 갱신(파일을 옮긴 경우), 같은 경로면 이름을
    바꾼다. 기본 템플릿이 없으면 이것이 기본이 된다. 반환: (항목, 새로 등록했는지)."""
    if not os.path.exists(path):
        raise FileNotFoundError(f"템플릿 없음: {path}")
    path = os.path.abspath(path)
    r = load_registry()
    name = (name or os.path.splitext(os.path.basename(path))[0]).strip()
    e = next((t for t in r["templates"] if t["name"] == name), None)
    if e is None:
        e = next((t for t in r["templates"] if _same_path(t["path"], path)), None)
        if e is not None and e["name"] != name:
            if r["default"] == e["name"]:
                r["default"] = name
            e["name"] = name
    new = e is None
    if new:
        e = {"name": name}
        r["templates"].append(e)
    e["path"] = path
    e["sha1"] = file_sha1(path)
    if make_default or not r["default"]:
        r["default"] = name
    save_registry(r)
    return e, new


def is_registered(path):
    return any(_same_path(t["path"], path) for t in load_registry()["templates"])


def resolve_template(ref):
    """--template 값 → 파일 경로. 있는 파일이면 그대로, 아니면 등록 이름으로 찾는다
    ('default'·'기본'·빈 값 = 기본 템플릿)."""
    if ref and os.path.exists(ref):
        return ref
    r = load_registry()
    name = r["default"] if (ref or "").strip() in DEFAULT_REFS else ref.strip()
    e = next((t for t in r["templates"] if t["name"] == name), None) if name else None
    if e is None:
        names = ", ".join(t["name"] for t in r["templates"]) or "없음"
        raise FileNotFoundError(f"'{ref}' — 파일도 아니고 등록된 템플릿 이름도 아닙니다(등록: {names})")
    if not os.path.exists(e["path"]):
        raise FileNotFoundError(f"등록된 템플릿 '{name}' 의 파일이 없습니다: {e['path']} — 옮겼다면 "
                                f"template_registry.py add <새 경로> --name {name}")
    return e["path"]


def describe(m):
    """검토용 요약(사람이 읽는 줄 목록)."""
    lay = m["layouts"]
    name = lambda r: f"'{r['name']}'" if r else "없음"
    out = [f"템플릿: {m['template']['name']} ({'세로' if m['size']['orientation'] == 'portrait' else '가로'} "
           f"{m['size']['w']}x{m['size']['h']}in)",
           f"레이아웃: 표지={name(lay['cover'])} · 목차={name(lay['toc'])} · 본문={name(lay['content'])}"
           f" · 간지=없음(본문 머리에 장 제목)"]
    label = {"system": "시스템명", "year": "연도", "meta": "부제(버전·날짜)", "audience": "대상",
             "kind": "매뉴얼 구분", "clear": "비움"}
    parts = []
    for s in m["cover"]["slots"]:
        where = f"자리표시자 {s['ph']}" if "ph" in s else f"고정 텍스트 '{s['static'][:14]}'"
        parts.append(f"{label.get(s['field'], s['field'])}←{where}")
    out.append("표지: " + " · ".join(parts))
    for role, spec in (("목차", m.get("toc") or {}), ("본문", m["content"])):
        rh = [s for s in spec.get("chrome", []) if s.get("role") == "run_head"]
        if rh:
            s = rh[0]
            out.append(f"{role} 러닝헤더: " + (f"자리표시자 {s['ph']}" if s.get("ph") is not None
                                          else f"고정 텍스트 '{s['static'][:24]}' → 이번 문구로 교체"))
    c = m["content"]
    out.append(f"본문 영역: 위 {c['body_top']}in · 아래 {c['body_bottom']}in"
               + (f" · 왼쪽 {c['body_x']}in" if c.get("body_x") is not None else "")
               + f" · 장/절 제목 원형 {'있음' if c.get('proto') else '없음(서식으로 새로 그림)'}")
    col = m["colors"]
    out.append(f"색: 강조 #{col['accent']} · 제목 #{col['dark']}" + (f" · 선 #{col['rule']}" if col.get("rule") else ""))
    if m.get("fonts"):
        out.append("글꼴: " + ", ".join(m["fonts"][:6]) + (" 외" if len(m["fonts"]) > 6 else ""))
    names = {"cover": "표지", "toc": "목차", "content": "본문"}
    for role, d in (m.get("decor") or {}).items():
        out.append(f"표본 {names.get(role, role)} 쪽 장식(선·도형·그림): {len(d['ids'])}개를 쪽마다 옮김"
                   + (f" · {d['skipped']}개는 본문·항목과 겹치거나 옮길 수 없어 뺌" if d.get("skipped") else ""))
    if m["cover"].get("keep_static"):
        out.append("그대로 두는 고정 텍스트: " + ", ".join(f"'{t[:30]}'" for t in m["cover"]["keep_static"]))
    if m.get("guard"):
        out.append("잔존 감시(산출물에 남으면 ERROR): " + ", ".join(f"'{g[:30]}'" for g in m["guard"]))
    return out


# ---------- 산출 --------------------------------------------------------------

class TemplateDeck:
    """템플릿 사본 위에 슬라이드를 만든다 — 원본 파일은 읽기만 한다."""

    def __init__(self, manifest):
        self.m = manifest
        self.prs = Presentation(manifest["template"]["path"])
        slides = list(self.prs.slides)
        # 장·절 제목 복제 원형은 표본 슬라이드를 지우기 전에 떠 둔다
        self.protos = {}
        ci = manifest["content"].get("proto_slide")
        if ci is not None and ci < len(slides):
            for key, spec in manifest["content"].get("proto", {}).items():
                sh = next((x for x in slides[ci].shapes if x.shape_id == spec["shape_id"]), None)
                if sh is not None:
                    self.protos[key] = copy.deepcopy(sh._element)
        self.decor = {}
        for role, d in (manifest.get("decor") or {}).items():
            if d["slide"] < len(slides):
                self.decor[role] = self._capture(slides[d["slide"]], set(d["ids"]))
        self._drop_slides()
        self.layouts = {}
        for role, ref in manifest["layouts"].items():
            if ref:
                self.layouts[role] = self._layout(ref)
        self.role_of = {id(lay.part): role for role, lay in self.layouts.items()}
        self.warnings = []

    def _fit(self, el, lay_el=None, lay_tree=None, master_el=None):
        w = fit_line(el, self.prs.slide_width, lay_el, lay_tree, master_defaults(lay_el, master_el))
        if w and w not in self.warnings:
            self.warnings.append(w)

    def _layout(self, ref):
        try:
            lay = self.prs.slide_masters[ref["master"]].slide_layouts[ref["index"]]
            if lay.name == ref["name"]:
                return lay
        except IndexError:
            pass
        for m in self.prs.slide_masters:
            for lay in m.slide_layouts:
                if lay.name == ref["name"]:
                    return lay
        raise ValueError(f"템플릿에 레이아웃 '{ref['name']}' 이 없습니다 — 템플릿이 바뀌었으면 "
                         "--template-refresh 로 다시 분석하세요")

    def _drop_slides(self):
        """표본 슬라이드를 모두 지운다 — 관계를 끊으면 저장 시 파일에서도 빠진다."""
        lst = self.prs.slides._sldIdLst
        for sld in list(lst):
            self.prs.part.drop_rel(sld.rId)
            lst.remove(sld)

    # -- 레이아웃 사본의 고정 텍스트 ---------------------------------------------
    def _replace_static(self, layout, match, new):
        n = 0
        for sh in iter_text_shapes(layout.shapes):
            if sh.is_placeholder:
                continue
            if _text(sh) == (match or "").replace("\x0b", "\n").strip():
                set_text_keep_style(sh.text_frame._txBody, new)
                self._fit(sh._element)
                n += 1
        return n

    def prepare(self, v):
        """이번 매뉴얼 문구(v: system·kind·audience·year·meta·run_head)로 레이아웃 사본의
        고정 텍스트를 바꾼다. 표지의 매뉴얼 구분·대상, 목차·본문의 러닝헤더가 대상이다."""
        for s in self.m["cover"]["slots"]:
            if "static" in s and s["field"] in v:
                self._replace_static(self.layouts["cover"], s["static"], v[s["field"]])
        for role in ("toc", "content"):
            spec = self.m.get(role) or {}
            if role not in self.layouts:
                continue
            for s in spec.get("chrome", []):
                if s.get("role") == "run_head" and s.get("static"):
                    self._replace_static(self.layouts[role], s["static"], v["run_head"])

    # -- 표본 장식 ---------------------------------------------------------------
    @staticmethod
    def _capture(slide, ids):
        """표본 슬라이드의 장식 도형 사본과 그 관계(그림·링크) — 슬라이드를 지우기 전에 떠 둔다."""
        out = []
        for sh in slide.shapes:
            if sh.shape_id not in ids:
                continue
            rels = {}
            for x in sh._element.iter():
                for k, v in x.attrib.items():
                    if k.startswith(_R_NS) and v in slide.part.rels:
                        r = slide.part.rels[v]
                        rels[v] = (r.reltype, r.target_ref if r.is_external else r.target_part, r.is_external)
            out.append((copy.deepcopy(sh._element), rels))
        return out

    def _add_decor(self, slide, role):
        """장식을 맨 아래 층에 얹는다 — 관계는 새 쪽에 다시 걸고, 이름표(mg-chrome-decor)로
        검증 게이트(캡처 규격·본문 분량)에서 빠지게 한다."""
        tree = slide.shapes._spTree
        nid = max([int(x.get("id")) for x in tree.iter(qn("p:cNvPr")) if x.get("id", "").isdigit()] + [1])
        pos = 2                                   # nvGrpSpPr·grpSpPr 바로 다음 = 가장 아래
        for el, rels in self.decor.get(role, []):
            new = copy.deepcopy(el)
            remap = {old: slide.part.relate_to(tgt, rt, is_external=ext) for old, (rt, tgt, ext) in rels.items()}
            for x in new.iter():
                for k, v in list(x.attrib.items()):
                    if k.startswith(_R_NS) and v in remap:
                        x.set(k, remap[v])
            for c in new.iter(qn("p:cNvPr")):
                nid += 1
                c.set("id", str(nid))
                c.set("name", "mg-chrome-decor")
            tree.insert(pos, new)
            pos += 1

    # -- 슬라이드 ---------------------------------------------------------------
    def new_slide(self, role):
        slide = self.prs.slides.add_slide(self.layouts[role])
        self._add_decor(slide, role)
        return slide

    def fill(self, slide, idx, text, name=None):
        """자리표시자에 글자를 넣는다 — 레이아웃 자리표시자의 서식(안내문 런 서식 포함)을
        그대로 복사해, 채운 글자가 템플릿 표본과 같은 모양이 되게 한다."""
        ph = next((p for p in slide.placeholders if p.placeholder_format.idx == idx), None)
        if ph is None or text in (None, ""):
            return None
        lay = next((p for p in slide.slide_layout.placeholders if p.placeholder_format.idx == idx), None)
        el = ph._element
        old = el.find(qn("p:txBody"))
        src = lay._element.find(qn("p:txBody")) if lay is not None else None
        if src is not None:
            new = copy.deepcopy(src)
            set_text_keep_style(new, text)
            if old is not None:
                el.replace(old, new)
            else:
                el.append(new)
        else:
            ph.text_frame.text = text
        self._fit(el, lay._element if lay is not None else None, slide.slide_layout.shapes._spTree,
                  slide.slide_layout.slide_master._element)
        ph.name = name or f"mg-chrome-ph{idx}"
        return ph

    def chrome(self, slide, key):
        """이 슬라이드 레이아웃의 머리·꼬리 자리(role=key) 목록."""
        role = self.role_of.get(id(slide.slide_layout.part))
        spec = self.m.get(role) or {}
        return [s for s in spec.get("chrome", []) if s.get("role") == key]

    def fill_role(self, slide, key, text):
        for s in self.chrome(slide, key):
            if s.get("ph") is not None:
                self.fill(slide, s["ph"], text)

    def add_proto(self, slide, key, text):
        """표본 슬라이드의 장·절 제목 상자를 복제해 글자만 바꿔 얹는다."""
        el = self.protos.get(key)
        if el is None:
            return None
        new = copy.deepcopy(el)
        set_text_keep_style(new.find(qn("p:txBody")), text)
        ids = [int(x.get("id")) for x in slide.shapes._spTree.iter(qn("p:cNvPr")) if x.get("id", "").isdigit()]
        c = new.find(".//p:cNvPr", NS)
        c.set("id", str(max(ids + [1]) + 1))
        c.set("name", "mg-chrome-title")
        slide.shapes._spTree.insert_element_before(new, "p:extLst")
        self._fit(new)
        return new

    def finalize(self, v):
        """마무리 — 비어 있는 자리표시자를 지우고(편집 화면에 표본 안내문이 보이지 않게),
        레이아웃·마스터에 남은 표본 문구를 걷어 낸다."""
        for slide in self.prs.slides:
            for ph in list(slide.placeholders):
                has_text = getattr(ph, "has_text_frame", False) and ph.text_frame.text.strip()
                if not has_text:
                    ph._element.getparent().remove(ph._element)
        guard = [g for g in self.m.get("guard", []) if g not in v.get("allow", "")]
        if not guard:
            return
        parts = [lay for m in self.prs.slide_masters for lay in m.slide_layouts] + list(self.prs.slide_masters)
        for part in parts:
            for sh in iter_text_shapes(part.shapes):
                for t in sh._element.iter(qn("a:t")):
                    for g in guard:
                        if t.text and g in t.text:
                            # 고정 텍스트는 이번 시스템명으로, 안내문은 중립 표기로(편집 화면 전용)
                            t.text = t.text.replace(g, "시스템명" if sh.is_placeholder else v.get("system", ""))


def leftover_check(pptx_path, guard, allow=""):
    """산출물(슬라이드·레이아웃·마스터·노트)에 템플릿 표본 문구가 남았는지 — 태그를 걷어 내고
    글자만 이어 붙여 보므로 런 경계로 쪼개진 문구도 잡는다. allow 에 든 문구(이번 매뉴얼이
    정말 그 이름을 쓰는 경우)는 제외한다. 반환: [(파일, 문구)]."""
    hits = []
    guard = [g for g in guard if g and g not in allow]
    if not guard:
        return hits
    with zipfile.ZipFile(pptx_path) as z:
        for n in z.namelist():
            if n.endswith(".xml") and n.startswith(("ppt/slides/", "ppt/slideLayouts/",
                                                     "ppt/slideMasters/", "ppt/notesSlides/")):
                txt = re.sub(r"<[^>]+>", "", z.read(n).decode("utf-8", "ignore"))
                for g in guard:
                    if g in txt:
                        hits.append((n, g))
    return hits
